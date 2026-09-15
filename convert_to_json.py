"""
Script chuyển đổi dữ liệu từ file Excel sang JSON.
Chạy 1 lần trên máy local để tạo file data.json.
"""
import pandas as pd
import json
import re
import os
import sys
import xlrd
import openpyxl

# Fix encoding on Windows
sys.stdout.reconfigure(encoding='utf-8')


EXAM_FILES = {
    "Kế toán trưởng, Trưởng-Phó phòng": "data_ktt.xls",
    "Chuyên viên": "data_cv.xls",
}

# Files for Khối VP – each file maps to a topic (display name)
VP_DIR = "Khối VP"
VP_FILES = [
    # (filename, sheet_name_or_None, icon, display_name)
    # sheet_name=None means read the first (or only) sheet
    ("5.SHNV Hanh chinh VP 2026.xls", "NV Hành Chính VP (100)", "📋", "Hành chính Văn phòng"),
    ("5.SHNV Hanh chinh VP 2026.xls", "AT Thông Tin (20)", "🔒", "An toàn Thông tin"),
    ("5.SHNV Hanh chinh VP 2026.xls", "VH Doanh Nghiệp (30)", "🏢", "Văn hóa Doanh nghiệp"),
    ("1.20CauATTT-TongQuan.xlsx", None, "🛡️", "ATTT - Tổng quan"),
    ("2.Vănhóadoanhnghiệp 2026 final.xls", None, "🌟", "Văn hóa DN (đề thi)"),
]

# Files for PCCL
PCCL_DIR = "PCCL"
PCCL_FILES = [
    # (filename, sheet_name_or_None, icon, display_name)
    ("20CauATTT-TongQuan.xlsx", None, "🛡️", "ATTT - Tổng quan"),
    ("Vănhóadoanhnghiệp 2026 final.xls", None, "🌟", "Văn hóa Doanh nghiệp"),
    ("NGB TCKT 2026 - 150c.xlsx", "Total", "📘", "Nâng giữ bậc TCKT 2026"),
]

# Đáp án chuẩn cho 150 câu NGB TCKT 2026 (1-indexed câu -> 0-indexed vị trí đáp án đúng)
NGB_ANSWERS = {
    1: 2, 2: 1, 3: 1, 4: 0, 5: 2, 6: 2, 7: 3, 8: 3, 9: 0, 10: 0,
    11: 0, 12: 2, 13: 3, 14: 3, 15: 0, 16: 1, 17: 1, 18: 1, 19: 1, 20: 1,
    21: 2, 22: 3, 23: 0, 24: 0, 25: 2, 26: 3, 27: 0, 28: 3, 29: 1, 30: 1,
    31: 3, 32: 2, 33: 3, 34: 0, 35: 3, 36: 3, 37: 3, 38: 0, 39: 3, 40: 0,
    41: 2, 42: 1, 43: 3, 44: 2, 45: 2, 46: 3, 47: 3, 48: 3, 49: 3, 50: 3,
    51: 2, 52: 2, 53: 3, 54: 1, 55: 3, 56: 2, 57: 1, 58: 1, 59: 2, 60: 0,
    61: 2, 62: 0, 63: 1, 64: 0, 65: 3, 66: 2, 67: 0, 68: 1, 69: 2, 70: 0,
    71: 1, 72: 0, 73: 1, 74: 0, 75: 2, 76: 3, 77: 1, 78: 3, 79: 1, 80: 0,
    81: 1, 82: 3, 83: 1, 84: 0, 85: 3, 86: 2, 87: 3, 88: 3, 89: 3, 90: 1,
    91: 3, 92: 3, 93: 2, 94: 2, 95: 2, 96: 2, 97: 0, 98: 2, 99: 2, 100: 0,
    101: 0, 102: 3, 103: 1, 104: 3, 105: 3, 106: 1, 107: 2, 108: 1, 109: 0, 110: 2,
    111: 0, 112: 2, 113: 0, 114: 0, 115: 1, 116: 3, 117: 2, 118: 2, 119: 2, 120: 1,
    121: 0, 122: 1, 123: 3, 124: 1, 125: 1, 126: 2, 127: 0, 128: 3, 129: 0, 130: 0,
    131: 0, 132: 0, 133: 0, 134: 0, 135: 2, 136: 1, 137: 1, 138: 2, 139: 3, 140: 0,
    141: 0, 142: 0, 143: 0, 144: 3, 145: 3, 146: 0, 147: 0, 148: 1, 149: 0, 150: 0,
}

TOPIC_INFO = {
    "KTT-130": ("📘", "KTT - 130 câu"),
    "CV-120": ("📗", "Chuyên viên - 120 câu"),
    "Tổng hợp": ("📚", "Tổng hợp"),
    "Thuế": ("💰", "Thuế"),
    "QC chi tieu noi bo": ("📋", "Quy chế chi tiêu nội bộ"),
    "Quản trị rủi ro": ("🛡️", "Quản trị rủi ro"),
    "ERP": ("💻", "ERP"),
    "kế toán": ("📊", "Kế toán"),
    "Chế độ kế toán- TT99": ("📑", "Chế độ kế toán - TT99"),
}


def parse_sheet(df):
    """Parse Q/A format (column 0 = 'Q' or 'A', column 1 = text, column 2 = 'X' for correct)."""
    questions = []
    current_q = None
    for _, row in df.iterrows():
        row_type = str(row.iloc[0]).strip() if pd.notna(row.iloc[0]) else ""
        text = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ""
        is_correct = str(row.iloc[2]).strip().upper() == "X" if pd.notna(row.iloc[2]) else False
        if row_type == "Q" and text:
            if current_q and current_q["answers"]:
                questions.append(current_q)
            current_q = {"question": text, "answers": [], "correct_idx": -1}
        elif row_type == "A" and text and current_q is not None:
            cleaned = re.sub(r'^[a-dA-D][\.\\)]\s*', '', text).strip()
            current_q["answers"].append(cleaned if cleaned else text)
            if is_correct:
                current_q["correct_idx"] = len(current_q["answers"]) - 1
    if current_q and current_q["answers"]:
        questions.append(current_q)
    return [q for q in questions if q["correct_idx"] >= 0]


def parse_sheet_stt(df):
    """Parse STT format (column 0 = number for question, a/b/c/d for answers,
    column 1 = text, column 2 = 'X' for correct answer).
    
    This format is used in some VP Excel sheets where:
    - Questions have a numeric STT (1.0, 2.0, ...) in column 0
    - Answers have a/b/c/d in column 0
    - Correct answer is marked with 'X' in column 2
    """
    questions = []
    current_q = None
    for _, row in df.iterrows():
        col0 = str(row.iloc[0]).strip() if pd.notna(row.iloc[0]) else ""
        text = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ""
        is_correct = str(row.iloc[2]).strip().upper() == "X" if len(row) > 2 and pd.notna(row.iloc[2]) else False

        # Skip header/empty rows
        if not col0 or not text:
            continue

        # Check if this row is a question (numeric STT like '1.0', '2.0', '10.0', etc.)
        is_question = False
        try:
            float(col0)
            is_question = True
        except ValueError:
            pass

        # Check if this row is an answer (a, b, c, d, etc.)
        is_answer = col0.lower() in ('a', 'b', 'c', 'd', 'e', 'f')

        if is_question and text:
            if current_q and current_q["answers"]:
                questions.append(current_q)
            # Clean question text: remove "Câu X:" prefix if present
            clean_text = re.sub(r'^Câu\s*\d+\s*:\s*', '', text).strip()
            current_q = {"question": clean_text, "answers": [], "correct_idx": -1}
        elif is_answer and text and current_q is not None:
            # Clean the answer text (remove trailing .0 for numeric answers)
            if text.endswith('.0') and text[:-2].isdigit():
                text = text[:-2]
            current_q["answers"].append(text)
            if is_correct:
                current_q["correct_idx"] = len(current_q["answers"]) - 1

    if current_q and current_q["answers"]:
        questions.append(current_q)
    return [q for q in questions if q["correct_idx"] >= 0]


def parse_sheet_ngb(filepath, answers_map=None):
    """Parse NGB TCKT format (bold question rows followed by 4 choice rows)."""
    if answers_map is None:
        answers_map = NGB_ANSWERS
    wb = openpyxl.load_workbook(filepath)
    s = wb['Total'] if 'Total' in wb.sheetnames else wb.active
    questions = []
    cur_q = None
    cur_opts = []

    for r in range(1, s.max_row + 1):
        c1 = s.cell(r, 1).value
        f = s.cell(r, 1).font
        if c1 is None:
            continue
        c1_str = str(c1).strip()
        if not c1_str:
            continue

        if f and f.b and r > 1:
            if cur_q and len(cur_opts) >= 2:
                q_idx = len(questions) + 1
                c_idx = answers_map.get(q_idx, 0)
                clean_q = re.sub(r'^Câu\s*\d+[\.:\s]*', '', cur_q).strip()
                questions.append({
                    "question": clean_q,
                    "answers": cur_opts,
                    "correct_idx": c_idx
                })
            cur_q = c1_str
            cur_opts = []
        elif cur_q is not None:
            cur_opts.append(c1_str)

    if cur_q and len(cur_opts) >= 2:
        q_idx = len(questions) + 1
        c_idx = answers_map.get(q_idx, 0)
        clean_q = re.sub(r'^Câu\s*\d+[\.:\s]*', '', cur_q).strip()
        questions.append({
            "question": clean_q,
            "answers": cur_opts,
            "correct_idx": c_idx
        })

    wb.close()
    return questions


def detect_format(df):
    """Detect whether the sheet uses Q/A format or STT (numeric) format.
    Returns 'qa' or 'stt'."""
    for _, row in df.iterrows():
        col0 = str(row.iloc[0]).strip() if pd.notna(row.iloc[0]) else ""
        if col0 == "Q":
            return "qa"
        if col0 in ("STT", "Phân Loại", "PHÂN LOẠI"):
            continue
        try:
            float(col0)
            return "stt"
        except ValueError:
            if col0.lower() in ('a', 'b', 'c', 'd'):
                return "stt"
    return "qa"  # default


def get_visible_sheets(filepath):
    visible_sheets = set()
    try:
        book = xlrd.open_workbook(filepath)
        for i in range(book.nsheets):
            if book.sheet_by_index(i).visibility == 0:
                visible_sheets.add(book.sheet_names()[i])
    except Exception:
        xls = pd.ExcelFile(filepath)
        visible_sheets = set(xls.sheet_names)
    return visible_sheets


def main():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    all_data = {}

    # ─── Existing exam files (Kế toán trưởng, Chuyên viên) ──────────
    for label, filename in EXAM_FILES.items():
        filepath = os.path.join(base_dir, filename)
        if not os.path.exists(filepath):
            print(f"❌ Không tìm thấy file: {filepath}")
            continue

        visible_sheets = get_visible_sheets(filepath)
        xls = pd.ExcelFile(filepath)
        topics = {}

        for sheet_name in xls.sheet_names:
            if sheet_name not in visible_sheets:
                continue
            df = pd.read_excel(xls, sheet_name=sheet_name, header=None)
            questions = parse_sheet(df)
            if questions:
                if label == "Chuyên viên" and sheet_name == "KTT-130":
                    continue
                if label == "Kế toán trưởng, Trưởng-Phó phòng" and sheet_name == "CV-120":
                    continue
                info = TOPIC_INFO.get(sheet_name, ("", sheet_name))
                display_name = f"{info[0]} {info[1]}"
                topics[display_name] = questions
                print(f"  ✅ {display_name}: {len(questions)} câu hỏi")

        all_data[label] = topics
        print(f"📦 {label}: {len(topics)} chủ đề")

    # ─── Khối VP files ───────────────────────────────────────────────
    vp_dir = os.path.join(base_dir, VP_DIR)
    vp_topics = {}

    if os.path.exists(vp_dir):
        print(f"\n--- Đang xử lý Khối VP ---")
        for filename, sheet_name, icon, display_name in VP_FILES:
            filepath = os.path.join(vp_dir, filename)
            if not os.path.exists(filepath):
                print(f"  ❌ Không tìm thấy file: {filename}")
                continue

            try:
                if sheet_name:
                    df = pd.read_excel(filepath, sheet_name=sheet_name, header=None)
                else:
                    df = pd.read_excel(filepath, header=None)

                # Detect format and parse
                fmt = detect_format(df)
                if fmt == "qa":
                    questions = parse_sheet(df)
                else:
                    questions = parse_sheet_stt(df)

                if questions:
                    full_name = f"{icon} {display_name}"
                    # Avoid duplicates if same topic from different files
                    if full_name in vp_topics:
                        full_name = f"{icon} {display_name} ({len(questions)} câu)"
                    vp_topics[full_name] = questions
                    print(f"  ✅ {full_name}: {len(questions)} câu hỏi (format: {fmt})")
                else:
                    print(f"  ⚠️ {display_name}: Không tìm thấy câu hỏi hợp lệ")
            except Exception as e:
                print(f"  ❌ Lỗi khi xử lý {filename} / {sheet_name}: {e}")
                import traceback
                traceback.print_exc()

        all_data["Khối VP"] = vp_topics
        print(f"📦 Khối VP: {len(vp_topics)} chủ đề")
    else:
        print(f"⚠️ Không tìm thấy thư mục: {vp_dir}")

    # ─── PCCL files ───────────────────────────────────────────────────
    pccl_dir = os.path.join(base_dir, PCCL_DIR)
    pccl_topics = {}

    if os.path.exists(pccl_dir):
        print(f"\n--- Đang xử lý PCCL ---")
        for filename, sheet_name, icon, display_name in PCCL_FILES:
            filepath = os.path.join(pccl_dir, filename)
            if not os.path.exists(filepath):
                print(f"  ❌ Không tìm thấy file: {filename}")
                continue

            try:
                if "150c" in filename:
                    questions = parse_sheet_ngb(filepath, NGB_ANSWERS)
                else:
                    if sheet_name:
                        df = pd.read_excel(filepath, sheet_name=sheet_name, header=None)
                    else:
                        df = pd.read_excel(filepath, header=None)

                    fmt = detect_format(df)
                    if fmt == "qa":
                        questions = parse_sheet(df)
                    else:
                        questions = parse_sheet_stt(df)

                if questions:
                    full_name = f"{icon} {display_name}"
                    pccl_topics[full_name] = questions
                    print(f"  ✅ {full_name}: {len(questions)} câu hỏi")
                else:
                    print(f"  ⚠️ {display_name}: Không tìm thấy câu hỏi hợp lệ")
            except Exception as e:
                print(f"  ❌ Lỗi khi xử lý {filename}: {e}")
                import traceback
                traceback.print_exc()

        all_data["PCCL"] = pccl_topics
        print(f"📦 PCCL: {len(pccl_topics)} chủ đề")
    else:
        print(f"⚠️ Không tìm thấy thư mục: {pccl_dir}")

    # ─── Write output ────────────────────────────────────────────────
    output_path = os.path.join(base_dir, "data.json")
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(all_data, f, ensure_ascii=False, indent=2)

    file_size = os.path.getsize(output_path) / 1024
    print(f"\n✅ Đã tạo data.json ({file_size:.1f} KB)")
    print(f"📍 Đường dẫn: {output_path}")


if __name__ == "__main__":
    main()
